"""Strict JSON/XLSX inputs for non-publishable EOD Social Studio previews."""
from __future__ import annotations

import io
import json
import math
from datetime import date, datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

from app.services.eod_narrative import validate_narrative
from app.services.eod_social_analysis import analyze_eod_contracts
from scheduler.market_utils import trade_date_ct


MAX_PREVIEW_INPUT_BYTES = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".json", ".xlsx"}
FIXTURE_DIR = Path(__file__).resolve().parents[2] / "fixtures" / "social"
BUILTIN_FIXTURES = {
    "positive": FIXTURE_DIR / "eod_preview_positive.json",
    "negative": FIXTURE_DIR / "eod_preview_negative.json",
}
CT = ZoneInfo("America/Chicago")

_TOP_LEVEL_FIELDS = {"mode", "metadata", "options", "analysis", "strike_gex", "narrative"}
_METADATA_FIELDS = {"symbol", "trade_date", "snapshot_at", "spot_price"}
_OPTION_REQUIRED = {"type", "strike", "expiry", "gamma", "iv", "open_interest"}
_OPTION_OPTIONAL = {"delta", "theta", "vega", "volume"}
_ANALYSIS_FIELDS = {
    "net_gex_b", "gamma_flip", "call_wall", "call_wall_gex_b",
    "put_wall", "put_wall_gex_b", "hot_zone", "hot_zone_gex_b",
}
_STRIKE_FIELDS = {"strike", "call_gex_b", "put_gex_b"}
_XLSX_SHEETS = {"Metadata", "Options", "Analysis", "StrikeGEX", "Narrative"}


def _strict_fields(value, expected: set[str], label: str, *, allow_empty=False) -> dict:
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be an object")
    if allow_empty and not value:
        return value
    missing = sorted(expected - set(value))
    unknown = sorted(set(value) - expected)
    if missing:
        raise ValueError(f"{label} is missing fields: {', '.join(missing)}")
    if unknown:
        raise ValueError(f"{label} contains unknown fields: {', '.join(unknown)}")
    return value


def _finite(value, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric") from exc
    if not math.isfinite(number):
        raise ValueError(f"{label} must be finite")
    return number


def _iso_date(value, label: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as exc:
            raise ValueError(f"{label} must be an ISO date") from exc
    raise ValueError(f"{label} must be an ISO date")


def _iso_datetime(value, label: str) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError(f"{label} must be an ISO datetime") from exc
    else:
        raise ValueError(f"{label} must be an ISO datetime")
    if parsed.tzinfo is None:
        raise ValueError(f"{label} must include a timezone offset")
    return parsed.astimezone(CT)


def _parse_metadata(data: dict) -> dict:
    data = _strict_fields(data, _METADATA_FIELDS, "metadata")
    symbol = str(data["symbol"]).strip().upper()
    if symbol != "SPX":
        raise ValueError("manual EOD preview inputs must use symbol SPX")
    trade_day = _iso_date(data["trade_date"], "metadata.trade_date")
    snapshot_at = _iso_datetime(data["snapshot_at"], "metadata.snapshot_at")
    if snapshot_at.date() != trade_day:
        raise ValueError("metadata.snapshot_at must fall on metadata.trade_date in Central Time")
    expiry_at = datetime.combine(trade_day, time(15, 0), tzinfo=CT)
    if snapshot_at >= expiry_at:
        raise ValueError("metadata.snapshot_at must be before the 3:00 PM CT SPX expiry")
    spot = _finite(data["spot_price"], "metadata.spot_price")
    if spot <= 0:
        raise ValueError("metadata.spot_price must be positive")
    return {
        "symbol": symbol,
        "trade_day": trade_day,
        "snapshot_at": snapshot_at,
        "spot_price": spot,
    }


def _parse_options(rows, trade_day: date) -> list[dict]:
    if not isinstance(rows, list) or not rows:
        raise ValueError("raw mode requires at least one options row")
    parsed: list[dict] = []
    option_types: set[str] = set()
    allowed = _OPTION_REQUIRED | _OPTION_OPTIONAL
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"options row {index} must be an object")
        missing = sorted(_OPTION_REQUIRED - set(row))
        unknown = sorted(set(row) - allowed)
        if missing:
            raise ValueError(f"options row {index} is missing fields: {', '.join(missing)}")
        if unknown:
            raise ValueError(f"options row {index} contains unknown fields: {', '.join(unknown)}")
        option_type = str(row["type"]).strip().upper()
        if option_type not in {"C", "P"}:
            raise ValueError(f"options row {index} type must be C or P")
        expiry = _iso_date(row["expiry"], f"options row {index} expiry")
        if expiry != trade_day:
            raise ValueError(f"options row {index} expiry must equal metadata.trade_date")
        strike = _finite(row["strike"], f"options row {index} strike")
        gamma = _finite(row["gamma"], f"options row {index} gamma")
        iv = _finite(row["iv"], f"options row {index} iv")
        oi_number = _finite(row["open_interest"], f"options row {index} open_interest")
        if strike <= 0 or gamma <= 0:
            raise ValueError(f"options row {index} strike and gamma must be positive")
        if not 0 < iv <= 5:
            raise ValueError(f"options row {index} iv must be in the range (0, 5]")
        if oi_number < 0 or not oi_number.is_integer():
            raise ValueError(f"options row {index} open_interest must be a nonnegative integer")
        parsed_row = {
            "type": option_type,
            "strike": strike,
            "expiry": expiry,
            "gamma": gamma,
            "iv": iv,
            "open_interest": int(oi_number),
        }
        for optional in _OPTION_OPTIONAL:
            value = row.get(optional, 0)
            number = _finite(value if value not in (None, "") else 0, f"options row {index} {optional}")
            if optional == "volume":
                if number < 0 or not number.is_integer():
                    raise ValueError(
                        f"options row {index} volume must be a nonnegative integer"
                    )
                parsed_row[optional] = int(number)
            else:
                parsed_row[optional] = number
        parsed.append(parsed_row)
        option_types.add(option_type)
    if option_types != {"C", "P"}:
        raise ValueError("raw mode requires both call and put contracts")
    return parsed


def _close(left: float, right: float) -> bool:
    return math.isclose(left, right, rel_tol=1e-7, abs_tol=1e-7)


def _metrics_analysis(metadata: dict, values: dict, strike_rows) -> tuple[dict, dict]:
    values = _strict_fields(values, _ANALYSIS_FIELDS, "analysis")
    if not isinstance(strike_rows, list) or not strike_rows:
        raise ValueError("metrics mode requires strike_gex rows")
    scalars = {key: _finite(values[key], f"analysis.{key}") for key in _ANALYSIS_FIELDS}
    for field in ("gamma_flip", "call_wall", "put_wall", "hot_zone"):
        if scalars[field] <= 0:
            raise ValueError(f"analysis.{field} must be positive")
    if scalars["call_wall_gex_b"] <= 0:
        raise ValueError("analysis.call_wall_gex_b must be positive")
    if scalars["put_wall_gex_b"] >= 0:
        raise ValueError("analysis.put_wall_gex_b must be negative")
    if scalars["hot_zone_gex_b"] >= 0:
        raise ValueError("analysis.hot_zone_gex_b must be negative")

    rows: list[dict] = []
    seen: set[float] = set()
    for index, row in enumerate(strike_rows, start=1):
        row = _strict_fields(row, _STRIKE_FIELDS, f"strike_gex row {index}")
        strike = _finite(row["strike"], f"strike_gex row {index} strike")
        call_b = _finite(row["call_gex_b"], f"strike_gex row {index} call_gex_b")
        put_b = _finite(row["put_gex_b"], f"strike_gex row {index} put_gex_b")
        if strike <= 0 or strike in seen:
            raise ValueError("strike_gex strikes must be positive and unique")
        if call_b < 0 or put_b > 0:
            raise ValueError(f"strike_gex row {index} must have nonnegative call and nonpositive put GEX")
        seen.add(strike)
        rows.append({"strike": strike, "call_gex": call_b * 1e9, "put_gex": put_b * 1e9})
    rows.sort(key=lambda item: item["strike"])

    by_strike = {row["strike"]: row for row in rows}
    for field in ("call_wall", "put_wall", "hot_zone"):
        if scalars[field] not in by_strike:
            raise ValueError(f"analysis.{field} must exist in strike_gex")
    computed_call = max(rows, key=lambda row: row["call_gex"])
    computed_put = min(rows, key=lambda row: row["put_gex"])
    computed_hot = min(rows, key=lambda row: row["call_gex"] + row["put_gex"])
    if computed_call["strike"] != scalars["call_wall"]:
        raise ValueError("analysis.call_wall is inconsistent with strike_gex")
    if computed_put["strike"] != scalars["put_wall"]:
        raise ValueError("analysis.put_wall is inconsistent with strike_gex")
    if computed_hot["strike"] != scalars["hot_zone"]:
        raise ValueError("analysis.hot_zone is inconsistent with strike_gex")
    expected = {
        "net_gex_b": sum((row["call_gex"] + row["put_gex"]) / 1e9 for row in rows),
        "call_wall_gex_b": computed_call["call_gex"] / 1e9,
        "put_wall_gex_b": computed_put["put_gex"] / 1e9,
        "hot_zone_gex_b": (computed_hot["call_gex"] + computed_hot["put_gex"]) / 1e9,
    }
    for field, expected_value in expected.items():
        if not _close(scalars[field], expected_value):
            raise ValueError(f"analysis.{field} is inconsistent with strike_gex")

    net_raw = scalars["net_gex_b"] * 1e9
    regime = "POSITIVE" if net_raw >= 0 else "NEGATIVE"
    dealer_position = "LONG" if net_raw >= 0 else "SHORT"
    source_trade_date = trade_date_ct(metadata["trade_day"])
    chart_snapshot = {
        "symbol": "SPX",
        "spot_price": metadata["spot_price"],
        "trade_date": source_trade_date,
        "created_at": metadata["snapshot_at"],
        "net_gex": net_raw,
        "gamma_flip": scalars["gamma_flip"],
        "call_wall": scalars["call_wall"],
        "put_wall": scalars["put_wall"],
        "gex_by_strike": rows,
    }
    analysis = {
        "source_snapshot_ref": None,
        "source_trade_date": source_trade_date,
        "snapshot_at": metadata["snapshot_at"],
        "expiry_at": datetime.combine(metadata["trade_day"], time(15, 0), tzinfo=CT),
        "spot_price": metadata["spot_price"],
        "zero_dte_contract_count": None,
        "regime": regime,
        "dealer_position": dealer_position,
        "net_gex_raw": net_raw,
        "net_gex_b": scalars["net_gex_b"],
        "gamma_flip": scalars["gamma_flip"],
        "call_wall": scalars["call_wall"],
        "call_wall_gex_raw": scalars["call_wall_gex_b"] * 1e9,
        "call_wall_gex_b": scalars["call_wall_gex_b"],
        "put_wall": scalars["put_wall"],
        "put_wall_gex_raw": scalars["put_wall_gex_b"] * 1e9,
        "put_wall_gex_b": scalars["put_wall_gex_b"],
        "hot_zone": scalars["hot_zone"],
        "hot_zone_gex_raw": scalars["hot_zone_gex_b"] * 1e9,
        "hot_zone_gex_b": scalars["hot_zone_gex_b"],
    }
    return chart_snapshot, analysis


def _normalise_payload(payload: dict, *, narrative_source: str) -> dict:
    payload = _strict_fields(payload, _TOP_LEVEL_FIELDS, "preview input")
    mode = str(payload["mode"]).strip().lower()
    if mode not in {"raw", "metrics"}:
        raise ValueError("mode must be raw or metrics")
    if narrative_source not in {"file", "lm_studio"}:
        raise ValueError("narrative_source must be file or lm_studio")
    metadata = _parse_metadata(payload["metadata"])
    if mode == "raw":
        if payload["analysis"] not in ({}, None):
            raise ValueError("raw mode analysis must be empty")
        if payload["strike_gex"] not in ([], None):
            raise ValueError("raw mode strike_gex must be empty")
        contracts = _parse_options(payload["options"], metadata["trade_day"])
        chart_snapshot, analysis = analyze_eod_contracts(
            contracts,
            metadata["spot_price"],
            metadata["trade_day"],
            metadata["snapshot_at"],
        )
    else:
        if payload["options"] not in ([], None):
            raise ValueError("metrics mode options must be empty")
        chart_snapshot, analysis = _metrics_analysis(
            metadata, payload["analysis"], payload["strike_gex"]
        )

    narrative = None
    if narrative_source == "file":
        narrative = validate_narrative(payload["narrative"])
    elif payload["narrative"] is not None and not isinstance(payload["narrative"], dict):
        raise ValueError("narrative must be an object when present")
    return {
        "mode": mode,
        "metadata": metadata,
        "chart_snapshot": chart_snapshot,
        "analysis": analysis,
        "narrative": narrative,
        "narrative_source": narrative_source,
    }


def _sheet_key_values(sheet, sheet_name: str) -> dict:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or tuple(str(value or "").strip().lower() for value in rows[0][:2]) != ("field", "value"):
        raise ValueError(f"{sheet_name} sheet must start with Field and Value headers")
    result: dict = {}
    for row_number, row in enumerate(rows[1:], start=2):
        key = str(row[0] or "").strip()
        value = row[1] if len(row) > 1 else None
        if not key and value in (None, ""):
            continue
        if not key or key in result:
            raise ValueError(f"{sheet_name} sheet row {row_number} has a missing or duplicate field")
        result[key] = value
    return result


def _sheet_records(
    sheet,
    sheet_name: str,
    *,
    required_headers: set[str],
    allowed_headers: set[str],
) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not header for header in headers) or len(headers) != len(set(headers)):
        raise ValueError(f"{sheet_name} sheet has invalid headers")
    missing = sorted(required_headers - set(headers))
    unknown = sorted(set(headers) - allowed_headers)
    if missing:
        raise ValueError(f"{sheet_name} sheet is missing columns: {', '.join(missing)}")
    if unknown:
        raise ValueError(f"{sheet_name} sheet contains unknown columns: {', '.join(unknown)}")
    result: list[dict] = []
    for row in rows[1:]:
        values = list(row[:len(headers)])
        if all(value in (None, "") for value in values):
            continue
        result.append(dict(zip(headers, values)))
    return result


def _parse_xlsx(content: bytes) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read EOD preview workbooks") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("uploaded file is not a readable .xlsx workbook") from exc
    try:
        actual = set(workbook.sheetnames)
        if actual != _XLSX_SHEETS:
            missing = sorted(_XLSX_SHEETS - actual)
            unknown = sorted(actual - _XLSX_SHEETS)
            details = []
            if missing:
                details.append(f"missing: {', '.join(missing)}")
            if unknown:
                details.append(f"unknown: {', '.join(unknown)}")
            raise ValueError(f"workbook sheets must be exactly {_XLSX_SHEETS} ({'; '.join(details)})")
        metadata = _sheet_key_values(workbook["Metadata"], "Metadata")
        mode = metadata.pop("mode", None)
        return {
            "mode": mode,
            "metadata": metadata,
            "options": _sheet_records(
                workbook["Options"],
                "Options",
                required_headers=_OPTION_REQUIRED,
                allowed_headers=_OPTION_REQUIRED | _OPTION_OPTIONAL,
            ),
            "analysis": _sheet_key_values(workbook["Analysis"], "Analysis"),
            "strike_gex": _sheet_records(
                workbook["StrikeGEX"],
                "StrikeGEX",
                required_headers=_STRIKE_FIELDS,
                allowed_headers=_STRIKE_FIELDS,
            ),
            "narrative": _sheet_key_values(workbook["Narrative"], "Narrative"),
        }
    finally:
        workbook.close()


def load_preview_bytes(
    content: bytes,
    filename: str,
    *,
    narrative_source: str = "file",
) -> dict:
    if not isinstance(content, bytes) or not content:
        raise ValueError("preview input file is empty")
    if len(content) > MAX_PREVIEW_INPUT_BYTES:
        raise ValueError("preview input file exceeds the 5 MB limit")
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("preview input must be a .json or .xlsx file")
    if suffix == ".json":
        try:
            payload = json.loads(content.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("uploaded JSON is invalid") from exc
    else:
        payload = _parse_xlsx(content)
    result = _normalise_payload(payload, narrative_source=narrative_source)
    result["input_format"] = suffix[1:]
    result["input_filename"] = Path(filename).name
    return result


def load_builtin_fixture(scenario: str, *, narrative_source: str = "file") -> dict:
    path = BUILTIN_FIXTURES.get(scenario)
    if path is None:
        raise ValueError("unknown EOD preview fixture")
    result = load_preview_bytes(path.read_bytes(), path.name, narrative_source=narrative_source)
    result["builtin_scenario"] = scenario
    return result
