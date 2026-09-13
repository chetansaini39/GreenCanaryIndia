import io
import json
import struct
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from app.models import post_templates
from app.services import social_generation
from app.services.eod_preview import (
    MAX_PREVIEW_INPUT_BYTES,
    load_builtin_fixture,
    load_preview_bytes,
)


USER_ID = "000000000000000000000001"


@pytest.fixture(autouse=True)
def _fixed_eod_theme(monkeypatch):
    monkeypatch.setattr(
        social_generation,
        "get_eod_theme",
        lambda db: {
            "brand_name": "GEX Intelligence",
            "brand_handle": "@gexintelligence",
            "background_color": "#0D1117",
            "surface_color": "#151A21",
            "text_color": "#E6EDF3",
            "muted_color": "#8B949E",
            "positive_color": "#22C55E",
            "negative_color": "#EF4444",
            "amber_color": "#D29922",
            "accent_color": "#1D9BF0",
            "spot_color": "#FFD700",
            "logo_path": None,
        },
    )


def _narrative():
    return {
        "regime_explanation": "Dealers may chase direction and amplify momentum",
        "hot_zone_explanation": "Breaks can run",
        "bull_behavior": "Upside can squeeze higher",
        "bear_behavior": "Dealer selling can accelerate",
        "wildcard_commentary": "Opening flow",
        "verdict": "Directional risk stays elevated until price reclaims the flip",
    }


def _raw_payload():
    return {
        "mode": "raw",
        "metadata": {
            "symbol": "SPX",
            "trade_date": "2026-08-26",
            "snapshot_at": "2026-08-26T14:00:00-05:00",
            "spot_price": 100,
        },
        "options": [
            {
                "type": "P", "strike": 99, "expiry": "2026-08-26",
                "gamma": 0.01, "iv": 0.25, "open_interest": 100,
            },
            {
                "type": "C", "strike": 101, "expiry": "2026-08-26",
                "gamma": 0.01, "iv": 0.25, "open_interest": 100,
            },
        ],
        "analysis": {},
        "strike_gex": [],
        "narrative": _narrative(),
    }


def _metrics_payload(scenario="negative"):
    return json.loads(
        (Path(__file__).resolve().parents[1] / "fixtures" / "social"
         / f"eod_preview_{scenario}.json").read_text()
    )


def _json_input(payload, name="preview.json", narrative_source="file"):
    return load_preview_bytes(
        json.dumps(payload).encode(), name, narrative_source=narrative_source
    )


def _xlsx_bytes(payload):
    workbook = Workbook()
    workbook.remove(workbook.active)

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Field", "Value"])
    metadata.append(["mode", payload["mode"]])
    for key, value in payload["metadata"].items():
        metadata.append([key, value])

    options = workbook.create_sheet("Options")
    option_headers = [
        "type", "strike", "expiry", "gamma", "iv", "open_interest",
        "delta", "theta", "vega", "volume",
    ]
    options.append(option_headers)
    for row in payload["options"]:
        options.append([row.get(key) for key in option_headers])

    analysis = workbook.create_sheet("Analysis")
    analysis.append(["Field", "Value"])
    for key, value in payload["analysis"].items():
        analysis.append([key, value])

    strikes = workbook.create_sheet("StrikeGEX")
    strike_headers = ["strike", "call_gex_b", "put_gex_b"]
    strikes.append(strike_headers)
    for row in payload["strike_gex"]:
        strikes.append([row.get(key) for key in strike_headers])

    narrative = workbook.create_sheet("Narrative")
    narrative.append(["Field", "Value"])
    for key, value in payload["narrative"].items():
        narrative.append([key, value])

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("payload", [_raw_payload(), _metrics_payload()])
def test_json_and_excel_inputs_are_equivalent(payload):
    json_result = _json_input(payload)
    excel_result = load_preview_bytes(_xlsx_bytes(payload), "preview.xlsx")

    assert excel_result["mode"] == json_result["mode"]
    assert excel_result["analysis"]["regime"] == json_result["analysis"]["regime"]
    assert excel_result["analysis"]["dealer_position"] == json_result["analysis"]["dealer_position"]
    assert excel_result["analysis"]["net_gex_b"] == pytest.approx(
        json_result["analysis"]["net_gex_b"]
    )
    assert excel_result["analysis"]["gamma_flip"] == pytest.approx(
        json_result["analysis"]["gamma_flip"]
    )
    assert excel_result["narrative"] == json_result["narrative"]


def test_raw_mode_runs_production_math_and_gamma_flip():
    result = _json_input(_raw_payload())
    analysis = result["analysis"]

    assert analysis["source_trade_date"].date() == date(2026, 8, 26)
    assert analysis["zero_dte_contract_count"] == 2
    assert analysis["net_gex_raw"] == pytest.approx(0)
    assert analysis["call_wall"] == 101
    assert analysis["put_wall"] == 99
    assert analysis["hot_zone"] == 99
    assert analysis["call_wall_gex_b"] == pytest.approx(0.00001)
    assert analysis["put_wall_gex_b"] == pytest.approx(-0.00001)
    assert 99 < analysis["gamma_flip"] < 101
    assert analysis["regime"] == "POSITIVE"


@pytest.mark.parametrize(
    "mutate,error",
    [
        (lambda p: p["metadata"].update({"unexpected": 1}), "unknown fields"),
        (lambda p: p["options"][0].update({"iv": 0}), "iv must be"),
        (lambda p: p["options"][0].update({"gamma": -0.1}), "gamma must be positive"),
        (lambda p: p["options"][0].update({"open_interest": -1}), "nonnegative integer"),
        (lambda p: p["options"][0].update({"expiry": "2026-08-27"}), "must equal"),
        (lambda p: p["options"][0].update({"expiry": "2026-08-26junk"}), "ISO date"),
        (lambda p: p.update({"analysis": {"net_gex_b": 1}}), "analysis must be empty"),
    ],
)
def test_raw_input_rejects_invalid_contract_data(mutate, error):
    payload = _raw_payload()
    mutate(payload)
    with pytest.raises(ValueError, match=error):
        _json_input(payload)


def test_metrics_mode_rejects_inconsistent_values():
    payload = _metrics_payload()
    payload["analysis"]["net_gex_b"] = 99
    with pytest.raises(ValueError, match="net_gex_b is inconsistent"):
        _json_input(payload)


def test_raw_mode_rejects_snapshot_after_expiry_and_absent_crossing():
    after_expiry = _raw_payload()
    after_expiry["metadata"]["snapshot_at"] = "2026-08-26T15:00:00-05:00"
    with pytest.raises(ValueError, match="before the 3:00 PM"):
        _json_input(after_expiry)

    no_crossing = _raw_payload()
    no_crossing["options"] = [
        {
            "type": "C", "strike": 100, "expiry": "2026-08-26",
            "gamma": 0.01, "iv": 0.25, "open_interest": 100,
        },
        {
            "type": "P", "strike": 100, "expiry": "2026-08-26",
            "gamma": 0.01, "iv": 0.25, "open_interest": 200,
        },
    ]
    with pytest.raises(ValueError, match="no reliable zero-gamma crossing"):
        _json_input(no_crossing)


def test_lm_studio_selection_ignores_file_narrative():
    payload = _metrics_payload()
    payload["narrative"] = {"untrusted": "ignored in favor of the model"}
    result = _json_input(payload, narrative_source="lm_studio")
    assert result["narrative"] is None
    assert result["narrative_source"] == "lm_studio"


def test_input_size_extension_and_workbook_sheets_are_strict():
    with pytest.raises(ValueError, match="5 MB"):
        load_preview_bytes(b"x" * (MAX_PREVIEW_INPUT_BYTES + 1), "preview.json")
    with pytest.raises(ValueError, match=".json or .xlsx"):
        load_preview_bytes(b"data", "preview.csv")

    workbook = Workbook()
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    with pytest.raises(ValueError, match="workbook sheets must be exactly"):
        load_preview_bytes(stream.getvalue(), "preview.xlsx")


def _fake_media_assets(tmp_path):
    assets = []
    for index, kind in enumerate(("chart", "regime", "levels", "playbook", "verdict")):
        path = (tmp_path / f"post-{index + 1}.png").resolve()
        path.write_bytes(
            b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\r" + b"IHDR"
            + struct.pack(">II", 1200, 675)
        )
        assets.append({
            "post_index": index,
            "kind": kind,
            "path": str(path),
            "alt_text": f"Preview image {index + 1}",
            "width": 1200,
            "height": 675,
            "size_bytes": path.stat().st_size,
        })
    return assets


def test_preview_generation_stores_five_posts_and_preview_boundary(monkeypatch, tmp_path):
    preview = load_builtin_fixture("negative")
    assets = _fake_media_assets(tmp_path)
    captured = {}
    market_analysis = MagicMock(side_effect=AssertionError("production market read"))
    monkeypatch.setattr(social_generation, "build_eod_analysis", market_analysis)
    monkeypatch.setattr(
        social_generation.post_templates,
        "get_template",
        lambda *args: {"template_text": post_templates._DEFAULTS["eod"]},
    )
    monkeypatch.setattr(social_generation, "render_eod_thread_media", lambda *args, **kwargs: assets)
    monkeypatch.setattr(
        social_generation.social_posts,
        "insert_draft",
        lambda db, doc: captured.update(doc) or doc["_id"],
    )

    post_id, chart_ok = social_generation.generate_eod_preview_draft(
        object(), preview, USER_ID
    )

    assert post_id == str(captured["_id"])
    assert chart_ok is True
    assert captured["trigger"] == "preview"
    assert captured["preview_only"] is True
    assert captured["source_snapshot_ref"] is None
    assert captured["input_mode"] == "metrics"
    assert captured["input_format"] == "json"
    assert captured["narrative_source"] == "file"
    assert captured["llm_model"] is None
    assert len(captured["tweets"]) == 5
    assert captured["media_assets"] == assets
    assert captured["media_version"] == 4
    assert captured["theme_snapshot"]["brand_handle"] == "@gexintelligence"
    assert captured["media_post_indexes"] == [0, 1, 2, 3, 4]
    assert captured["validation_result"]["preview_only"] is True
    market_analysis.assert_not_called()


def test_preview_lm_path_uses_existing_bounded_wrapper(monkeypatch, tmp_path):
    preview = load_builtin_fixture("positive", narrative_source="lm_studio")
    assets = _fake_media_assets(tmp_path)
    captured = {}
    generated = MagicMock(return_value=(_narrative(), "local-test-model"))
    monkeypatch.setattr(social_generation, "generate_eod_narrative", generated)
    monkeypatch.setattr(
        social_generation.post_templates,
        "get_template",
        lambda *args: {"template_text": post_templates._DEFAULTS["eod"]},
    )
    monkeypatch.setattr(social_generation, "render_eod_thread_media", lambda *args, **kwargs: assets)
    monkeypatch.setattr(
        social_generation.social_posts,
        "insert_draft",
        lambda db, doc: captured.update(doc) or doc["_id"],
    )

    social_generation.generate_eod_preview_draft(object(), preview, USER_ID)
    generated.assert_called_once_with(preview["analysis"])
    assert captured["llm_model"] == "local-test-model"
    assert captured["narrative_source"] == "lm_studio"


def test_preview_lm_failure_creates_no_draft_or_media(monkeypatch):
    preview = load_builtin_fixture("positive", narrative_source="lm_studio")
    insert = MagicMock()
    render = MagicMock()
    monkeypatch.setattr(
        social_generation.post_templates,
        "get_template",
        lambda *args: {"template_text": post_templates._DEFAULTS["eod"]},
    )
    monkeypatch.setattr(
        social_generation,
        "generate_eod_narrative",
        MagicMock(side_effect=RuntimeError("model unavailable")),
    )
    monkeypatch.setattr(social_generation, "render_eod_thread_media", render)
    monkeypatch.setattr(social_generation.social_posts, "insert_draft", insert)

    with pytest.raises(RuntimeError, match="model unavailable"):
        social_generation.generate_eod_preview_draft(object(), preview, USER_ID)
    render.assert_not_called()
    insert.assert_not_called()


def test_preview_publish_is_blocked_before_credentials(monkeypatch):
    monkeypatch.delenv("TWITTER_API_KEY", raising=False)
    monkeypatch.setattr(
        social_generation.social_posts,
        "find_by_id",
        lambda *args: {"status": "draft", "preview_only": True},
    )
    with pytest.raises(RuntimeError, match="Preview-only drafts cannot be published"):
        social_generation.publish_post(object(), "64b64b64b64b64b64b64b64b")
